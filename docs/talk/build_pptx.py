#!/usr/bin/env python3
"""Export the HTML talk to PPTX for Google Slides or PowerPoint, keeping the click-by-click builds.

Headless Chrome renders every build state of every slide at 1920x1080 (the deck's ?export&step=k#n view). Each PPTX
slide stacks its states as full-slide pictures, and every state after the first fades in on click, so a slide builds
the same way it does in the HTML deck. Speaker notes go into the notes pane. Slide text is part of the pictures, so it
is not editable in the export.

Needs python-pptx (`pip install python-pptx`) and Google Chrome.

Usage: build_pptx.py --html tpu_distil_talk.html --out tpu_distil_talk.pptx [--chrome google-chrome] [--jobs 4]
"""
from __future__ import annotations

import argparse
import html
import re
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)   # 16:9, 13.333 x 7.5 in
FADE_MS = 300
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def read_deck(path: Path):
  """[(build steps, speaker notes)] for every slide in the HTML deck."""
  page = path.read_text()
  secs = re.findall(r'<section class="deck-slide" id="s\d+">(.*?)<aside class="notes">(.*?)</aside></section>', page, re.S)
  return [(max([int(x) for x in re.findall(r'class="[^"]*\bbuild\b[^"]*" data-step="(\d+)"', body)], default=0), html.unescape(notes))
          for body, notes in secs]


def render(chrome: str, deck: Path, n: int, step: int, out: Path) -> Path:
  url = f"{deck.resolve().as_uri()}?export&step={step}#{n}"
  subprocess.run([chrome, "--headless=new", "--no-sandbox", "--disable-gpu", "--hide-scrollbars", "--window-size=1920,1080",
                  "--virtual-time-budget=3000", f"--screenshot={out}", url], check=True, capture_output=True, timeout=120)
  if not out.exists() or out.stat().st_size == 0:
    raise RuntimeError(f"no screenshot for slide {n} step {step}")
  return out


def click_fades(spids: list[int]) -> etree._Element:
  """<p:timing> that fades in each shape in `spids` on its own click, in order."""
  ids = iter(range(1, 10_000))
  clicks = []
  for spid in spids:
    clicks.append(f'''<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{next(ids)}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">
<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>
<p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" dur="{FADE_MS}"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>''')
  root, seq = next(ids), next(ids)
  xml = f'''<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="{root}" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="{seq}" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(clicks)}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
  return etree.fromstring(xml)


def main() -> None:
  ap = argparse.ArgumentParser()
  ap.add_argument("--html", type=Path, required=True)
  ap.add_argument("--out", type=Path, required=True)
  ap.add_argument("--chrome", default="google-chrome")
  ap.add_argument("--jobs", type=int, default=4)
  args = ap.parse_args()
  deck = read_deck(args.html)
  with tempfile.TemporaryDirectory() as tmp:
    jobs = [(n, k, Path(tmp) / f"s{n:03d}_{k}.png") for n, (steps, _) in enumerate(deck, 1) for k in range(steps + 1)]
    with ThreadPoolExecutor(args.jobs) as pool:
      list(pool.map(lambda j: render(args.chrome, args.html, *j), jobs))
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    for n, (steps, notes) in enumerate(deck, 1):
      slide = prs.slides.add_slide(blank)
      pics = [slide.shapes.add_picture(str(Path(tmp) / f"s{n:03d}_{k}.png"), 0, 0, SLIDE_W, SLIDE_H) for k in range(steps + 1)]
      for k, pic in enumerate(pics):
        pic.name = f"Slide {n} build {k}"
      if steps:
        slide._element.append(click_fades([pic.shape_id for pic in pics[1:]]))
      slide.notes_slide.notes_text_frame.text = notes.strip()
    prs.save(args.out)
  print(f"{len(deck)} slides, {len(jobs)} pictures, {sum(s for s, _ in deck)} click builds -> {args.out}")


if __name__ == "__main__":
  main()
